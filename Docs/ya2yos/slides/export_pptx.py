#!/usr/bin/env python3
"""Render the Typst defense slides into a visually equivalent PPTX file."""

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


SLIDES_DIR = Path(__file__).resolve().parent
REPO_ROOT = SLIDES_DIR.parents[2]
DEFAULT_SOURCE = SLIDES_DIR / "defense.typ"
DEFAULT_OUTPUT = SLIDES_DIR / "ya2yos-defense.pptx"


def page_number(path: Path) -> int:
    match = re.search(r"-(\d+)\.png$", path.name)
    if match is None:
        raise ValueError(f"unexpected rendered page name: {path.name}")
    return int(match.group(1))


def render_pdf(source: Path, destination: Path, dpi: int) -> list[Path]:
    with tempfile.TemporaryDirectory(prefix="ya2yos-defense-") as temp_dir:
        temp_path = Path(temp_dir)
        pdf = temp_path / "defense.pdf"
        prefix = temp_path / "slide"
        subprocess.run(
            ["typst", "compile", "--root", str(REPO_ROOT), str(source), str(pdf)],
            check=True,
            cwd=REPO_ROOT,
        )
        subprocess.run(
            ["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(prefix)],
            check=True,
        )
        pages = sorted(temp_path.glob("slide-*.png"), key=page_number)
        if not pages:
            raise RuntimeError("Typst PDF did not produce any rendered slide image")

        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = "Ya2yOS 内核设计与工程实践答辩"
        presentation.core_properties.subject = "由 Docs/ya2yos/slides/defense.typ 导出"
        presentation.core_properties.author = "Ya2yOS"

        blank_layout = presentation.slide_layouts[6]
        for page in pages:
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(page),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        presentation.save(destination)
        return pages


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--dpi", type=int, default=180)
    args = parser.parse_args()

    pages = render_pdf(args.source.resolve(), args.output.resolve(), args.dpi)
    print(f"created {args.output} with {len(pages)} slides at {args.dpi} DPI")


if __name__ == "__main__":
    main()
